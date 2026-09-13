"""Generate the public/client catalog and authoritative card database from the design workbook.

Usage: python scripts/import-card-catalog.py <workbook.xlsx>
The generated TypeScript files are committed, so production builds never depend on Excel.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
FACTIONS = {"鱼人": "MERMAID", "巨人": "GIANT", "通用": "COMMON"}
TYPES = {"随从": "MINION", "法术": "SPELL"}
RARITIES = {"普通": "COMMON", "稀有": "RARE", "史诗": "EPIC", "传说": "LEGENDARY"}
RUNES = {"MERMAID": "◉", "GIANT": "◆", "COMMON": "✦"}

# v0.4.x only introduces the formal pool and faction rules. Effects not yet supported by
# the small MVP engine remain playable prototype spells while retaining their exact text.
SUPPORTED_EFFECTS = {
    "GEN-S-002": '{ type: "DEAL_DAMAGE", amount: 1, target: "ANY_ENEMY" }',
    "GEN-S-003": '{ type: "HEAL", amount: 3, target: "FRIENDLY_HERO" }',
    "GEN-S-004": '{ type: "BUFF", attack: 2, health: 0, target: "FRIENDLY_MINION" }',
    "GEN-S-005": '{ type: "DRAW_CARD", amount: 1 }',
    "GEN-S-006": '{ type: "DEAL_DAMAGE", amount: 2, target: "ANY_ENEMY" }',
    "GEN-S-008": '{ type: "BUFF", attack: 0, health: 3, target: "FRIENDLY_MINION" }',
    "GEN-S-009": '{ type: "DRAW_CARD", amount: 2 }',
    "GEN-S-011": '{ type: "HEAL", amount: 5, target: "FRIENDLY_HERO" }',
    "GEN-S-012": '{ type: "BUFF", attack: 2, health: 1, target: "FRIENDLY_MINION" }',
    "GEN-S-014": '{ type: "DEAL_DAMAGE", amount: 5, target: "ANY_ENEMY" }',
    "GEN-S-017": '{ type: "HEAL", amount: 3, target: "FRIENDLY_HERO" }',
    "GEN-S-019": '{ type: "DRAW_CARD", amount: 3 }',
    "GEN-S-020": '{ type: "HEAL", amount: 4, target: "FRIENDLY_HERO" }',
}


def js(value: object) -> str:
    return json.dumps(value, ensure_ascii=False)


def normalized_keywords(row: dict[str, object]) -> list[str]:
    text = f"{row['关键词']} {row['效果描述']} {row['实现标签']}"
    mapping = [
        ("伤害", "DAMAGE"), ("治疗", "HEAL"), ("抽", "DRAW"), ("检索", "DRAW"),
        ("增益", "BUFF"), ("守护", "GUARD"), ("突袭", "CHARGE"), ("冲锋", "CHARGE"),
        ("汲取", "DRAIN"), ("护盾", "BARRIER"),
    ]
    return list(dict.fromkeys(code for marker, code in mapping if marker in text))


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("usage: import-card-catalog.py <workbook.xlsx>")
    workbook = load_workbook(Path(sys.argv[1]), data_only=False)
    sheet = workbook["卡牌总表"]
    headers = [cell.value for cell in sheet[8]]
    rows: list[dict[str, object]] = []
    for values in sheet.iter_rows(min_row=9, values_only=True):
        if not values[0]:
            continue
        row = dict(zip(headers, values))
        if row.get("卡牌ID") and row.get("阵营") in FACTIONS:
            rows.append(row)
    if len(rows) != 60:
        raise ValueError(f"expected 60 cards, found {len(rows)}")

    db_lines = [
        'import type { CardDefinition } from "@riftbound/shared";', "",
        "// Generated from Mermaids-Giants-Terra_首批60张卡牌策划表.xlsx.",
        "// Complex effects tagged PROTOTYPE retain their formal text but await a later rules-engine version.",
        "export const CARD_DATABASE: readonly CardDefinition[] = [",
    ]
    catalog_lines = [
        'import type { CardFaction, CardRarity } from "@riftbound/shared";', "",
        'export type { CardRarity };',
        'export type CatalogCardType = "MINION" | "SPELL";',
        'export type CardKeyword = "DAMAGE" | "HEAL" | "DRAW" | "BUFF" | "GUARD" | "CHARGE" | "DRAIN" | "BARRIER";', "",
        "export interface CatalogCard {", "  id: string;", "  nameKey: string;", "  descriptionKey: string;",
        "  type: CatalogCardType;", "  faction: CardFaction;", "  species?: string;", "  cost: number;",
        "  rarity: CardRarity;", "  rune: string;", "  keywords: readonly CardKeyword[];", "  deckLimit: number;",
        "  attack?: number;", "  health?: number;", "}", "",
        "// Public catalog only; it never contains match hand/deck order or any other secret state.",
        "export const CARD_CATALOG: readonly CatalogCard[] = [",
    ]
    translations: list[str] = []

    for row in rows:
        card_id = str(row["卡牌ID"])
        key = card_id.lower().replace("-", "_")
        faction = FACTIONS[str(row["阵营"])]
        card_type = TYPES[str(row["卡牌类型"])]
        rarity = RARITIES[str(row["稀有度"])]
        cost = int(row["费用"])
        limit = int(row["建议携带上限"])
        species = None if str(row["随从种族"]) == "n.a." else str(row["随从种族"])
        keywords = normalized_keywords(row)
        tags = [tag for tag in str(row["实现标签"]).split(";") if tag]
        common = (
            f'id: {js(card_id)}, nameKey: {js(key + "_name")}, descriptionKey: {js(key + "_description")}, '
            f'type: {js(card_type)}, faction: {js(faction)}, rarity: {js(rarity)}, cost: {cost}, '
            f'rune: {js(RUNES[faction])}, deckLimit: {limit}, keywords: {js(keywords)}, '
            f'targetRule: {js(str(row["目标规则"]))}, implementationTags: {js(tags)}'
        )
        if species:
            common += f", species: {js(species)}"
        if card_type == "MINION":
            definition = f"{{ {common}, attack: {int(row['攻击'])}, health: {int(row['生命'])} }}"
        else:
            definition = f"{{ {common}, effect: {SUPPORTED_EFFECTS.get(card_id, '{ type: \"PROTOTYPE\" }')} }}"
        db_lines.append(f"  {definition},")

        public = (
            f'{{ id: {js(card_id)}, nameKey: {js(key + "_name")}, descriptionKey: {js(key + "_description")}, '
            f'type: {js(card_type)}, faction: {js(faction)}, cost: {cost}, rarity: {js(rarity)}, '
            f'rune: {js(RUNES[faction])}, keywords: {js(keywords)}, deckLimit: {limit}'
        )
        if species:
            public += f", species: {js(species)}"
        if card_type == "MINION":
            public += f", attack: {int(row['攻击'])}, health: {int(row['生命'])}"
        catalog_lines.append(f"  {public} }},")
        translations.append(f"  {key}_name: {js(str(row['中文名']))},")
        translations.append(f"  {key}_description: {js(str(row['效果描述']))},")

    db_lines += [
        "];", "",
        'export const SPECIAL_BITCOIN_COIN_ID = "SPECIAL_BITCOIN_COIN";',
        'export const SPECIAL_BITCOIN_COIN: CardDefinition = {',
        '  id: SPECIAL_BITCOIN_COIN_ID,',
        '  nameKey: "special_bitcoin_coin_name",',
        '  descriptionKey: "special_bitcoin_coin_description",',
        '  type: "SPELL", faction: "COMMON", rarity: "LEGENDARY", cost: 0, rune: "₿", deckLimit: 0,',
        '  keywords: [], targetRule: "无需目标",',
        '  implementationTags: ["special", "second_player_bonus", "temporary_mana"],',
        '  effect: { type: "GAIN_TEMP_MANA", amount: 1 }',
        '};', "",
        "export const CARD_DEFINITIONS = new Map([...CARD_DATABASE, SPECIAL_BITCOIN_COIN].map((card) => [card.id, card]));",
        'const DEFAULT_COUNTS: Readonly<Record<string, number>> = {',
        '  "MER-M-001": 2, "MER-M-002": 2, "MER-M-003": 2, "MER-M-004": 2, "MER-M-005": 2,',
        '  "MER-M-006": 2, "MER-M-007": 2, "MER-M-008": 2, "MER-M-009": 2, "MER-M-010": 2,',
        '  "MER-M-011": 2, "MER-M-012": 2, "MER-M-013": 2, "MER-M-014": 2, "GEN-S-001": 2',
        '};',
        'export const DEFAULT_DECK_DEFINITION_IDS = Object.entries(DEFAULT_COUNTS).flatMap(([id, count]) => Array.from({ length: count }, () => id));', "",
        "export function getCardDefinition(definitionId: string): CardDefinition {",
        "  const definition = CARD_DEFINITIONS.get(definitionId);",
        "  if (!definition) throw new Error(`Unknown card definition: ${definitionId}`);",
        "  return definition;", "}", "",
    ]
    catalog_lines += [
        "];", "", 'export const RARITY_LABEL: Readonly<Record<CardRarity, string>> = {',
        '  COMMON: "普通", RARE: "稀有", EPIC: "史诗", LEGENDARY: "传说"', "};", "",
        'export const FACTION_LABEL: Readonly<Record<CardFaction, string>> = {',
        '  MERMAID: "鱼人", GIANT: "巨人", COMMON: "通用"', "};", "",
        'export const KEYWORD_LABEL: Readonly<Record<CardKeyword, string>> = {',
        '  DAMAGE: "伤害", HEAL: "治疗", DRAW: "抽牌", BUFF: "强化", GUARD: "守护",',
        '  CHARGE: "突袭", DRAIN: "汲取", BARRIER: "护盾"', "};", "",
    ]
    i18n_lines = [
        "export const ZH_CN: Readonly<Record<string, string>> = {",
        '  special_bitcoin_coin_name: "比特币",',
        '  special_bitcoin_coin_description: "本回合获得 +1 临时法力。",',
        *translations, "};", "",
        "export function translateZhCn(key: string): string {", "  return ZH_CN[key] ?? key;", "}", "",
    ]

    (ROOT / "packages/game-core/src/cards/database.ts").write_text("\n".join(db_lines), encoding="utf-8")
    (ROOT / "client/src/data/cardCatalog.ts").write_text("\n".join(catalog_lines), encoding="utf-8")
    (ROOT / "packages/shared/src/i18n/zh-CN.ts").write_text("\n".join(i18n_lines), encoding="utf-8")
    print("generated 60 cards: 20 MERMAID, 20 GIANT, 20 COMMON spells")


if __name__ == "__main__":
    main()
